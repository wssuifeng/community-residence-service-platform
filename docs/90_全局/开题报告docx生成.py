# -*- coding: utf-8 -*-
"""生成开题报告及任务书（v2，2026-10-22 版式重做）

版式基准取自学院样例（开题报告及任务书_样例1）：
  正文 = 自定义样式「样式1」，宋体 12pt（sz=24），首行缩进 2 字符（480 twips）
  标题 = 表头段落「Table Paragraph」，宋体 12pt，加粗，取消首行缩进
参考文献条目 = 样式1 + 取消首行缩进（长条目折行不缩进，便于阅读）
"""
import copy
import re

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm

BASE = 'D:/code/community-residence-service-platform/docs/90_全局'
SRC = BASE + '/参考资料/开题报告/开题报告及任务书_样例1（博客网站）.docx'
DRAFT = BASE + '/开题报告草稿.md'
# 输出到论文编写目录（docs/90_全局/论文编写）
# 输出路径；目标文件被 WPS/Word 占用时自动改写到同目录的 _new 文件
OUT = BASE + '/论文编写/开题报告及任务书_小区居住服务管理系统.docx'


def _writable(path):
    """探测目标文件是否可写（被 WPS/Word 打开时不可写）。"""
    import os
    if not os.path.exists(path):
        return True
    try:
        with open(path, 'ab'):
            return True
    except OSError:
        return False

RED = 'FF0000'
BLUE = '0070C0'

# ---------- 表头与个人信息（2026-10-22 用户提供） ----------
COLLEGE = '软件工程学院'
CLASS_NAME = '25Z250102'
STUDENT_NAME = '余荣艺'
STUDENT_NO = '2025101010276'
ADVISOR = '印丽娅'
ENTERPRISE_TEACHER = '暂无'
TITLE = '基于SpringBoot的小区居住服务管理系统设计与实现'


def extract_refs():
    """从草稿的 §4.1 主表读取参考文献条目（改表即改 docx）。"""
    txt = open(DRAFT, encoding='utf-8').read()
    start = txt.index('| # | 文献条目')
    # 主表止于下一个二级标题或说明段（兼容有无「备选池」两种写法）
    for marker in ('**备选池**', '### 4.2', '**说明**'):
        pos = txt.find(marker, start)
        if pos > 0:
            end = pos
            break
    else:
        end = len(txt)
    block = txt[start:end]
    rows = [l for l in block.split('\n') if re.match(r'\|\s*\d+\s*\|', l)]
    refs = []
    for line in rows:
        cells = line.split('|')
        refs.append(cells[2].strip().replace('**', ''))
    return refs


# ---------- 底层写入工具 ----------
def _ensure_rpr(run_el):
    rpr = run_el.find(qn('w:rPr'))
    if rpr is None:
        rpr = OxmlElement('w:rPr')
        run_el.insert(0, rpr)
    return rpr


def _set_run(run_el, text, bold=False, color=None, font='宋体', size=24):
    """写一个 run：宋体 + 12pt(sz=24)，可选加粗/颜色。"""
    for child in list(run_el):
        run_el.remove(child)
    rpr = OxmlElement('w:rPr')
    fonts = OxmlElement('w:rFonts')
    fonts.set(qn('w:ascii'), font)
    fonts.set(qn('w:hAnsi'), font)
    fonts.set(qn('w:eastAsia'), font)
    rpr.append(fonts)
    if bold:
        rpr.append(OxmlElement('w:b'))
    if color:
        c = OxmlElement('w:color')
        c.set(qn('w:val'), color)
        rpr.append(c)
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), str(size))
    rpr.append(sz)
    szcs = OxmlElement('w:szCs')
    szcs.set(qn('w:val'), str(size))
    rpr.append(szcs)
    run_el.append(rpr)
    t = OxmlElement('w:t')
    t.set(qn('xml:space'), 'preserve')
    t.text = text
    run_el.append(t)


def _new_par(style_id='8', indent_chars=200, before=2, align=None):
    """新建段落：style_id '8' = 模板的「样式1」（宋体 12pt，首行缩进 2 字符）。"""
    p = OxmlElement('w:p')
    ppr = OxmlElement('w:pPr')
    pstyle = OxmlElement('w:pStyle')
    pstyle.set(qn('w:val'), style_id)
    ppr.append(pstyle)
    if align:
        jc = OxmlElement('w:jc')
        jc.set(qn('w:val'), align)
        ppr.append(jc)
    spacing = OxmlElement('w:spacing')
    spacing.set(qn('w:before'), str(before))
    ppr.append(spacing)
    if indent_chars:
        ind = OxmlElement('w:ind')
        ind.set(qn('w:firstLine'), str(int(indent_chars * 2.4)))
        ind.set(qn('w:firstLineChars'), str(indent_chars))
        ppr.append(ind)
    p.append(ppr)
    return p


def clear_cell(cell):
    for p in list(cell._tc.findall(qn('w:p'))):
        cell._tc.remove(p)


def fill_cell(cell, lines):
    """lines: [(text, kind)]；text 可为字符串或字符串列表（列表按多段输出）。

    kind ∈ title / h2 / h3 / outline / bold / body / ref / red / blue
    """
    clear_cell(cell)
    for text, kind in lines:
        if isinstance(text, list):
            for one in text:
                _emit_par(cell, one, kind)
        else:
            _emit_par(cell, text, kind)


def _emit_par(cell, text, kind):
    if kind == 'title':
        p = _new_par(indent_chars=0, before=2)
        r = OxmlElement('w:r')
        _set_run(r, text, bold=True)
    elif kind == 'title_break':
        # 大章节另起一页：段首插入分页符，再接标题
        p = _new_par(indent_chars=0, before=2)
        pb = OxmlElement('w:r')
        br = OxmlElement('w:br')
        br.set(qn('w:type'), 'page')
        pb.append(br)
        p.append(pb)
        r = OxmlElement('w:r')
        _set_run(r, text, bold=True)
    elif kind == 'h2':
        # 节标题：顶格加粗，段前留白
        p = _new_par(indent_chars=0, before=6)
        r = OxmlElement('w:r')
        _set_run(r, text, bold=True)
    elif kind == 'h3':
        # 条标题：缩进 2 字符加粗
        p = _new_par(indent_chars=200, before=4)
        r = OxmlElement('w:r')
        _set_run(r, text, bold=True)
    elif kind == 'outline':
        # 论文提纲 / 进度表：按前导全角空格层级缩进
        lead = len(text) - len(text.lstrip('　'))
        p = _new_par(indent_chars=lead * 200 if lead else 0, before=0)
        r = OxmlElement('w:r')
        _set_run(r, text.lstrip('　'), size=21)
    elif kind == 'bold':
        p = _new_par(indent_chars=200, before=2)
        r = OxmlElement('w:r')
        _set_run(r, text, bold=True)
    elif kind == 'ref':
        # 参考文献：不首行缩进，长条目折行后对齐左边界
        p = _new_par(indent_chars=0, before=0)
        r = OxmlElement('w:r')
        _set_run(r, text, size=21)
    elif kind == 'red':
        p = _new_par(indent_chars=200, before=2)
        r = OxmlElement('w:r')
        _set_run(r, text, color=RED)
    elif kind == 'blue':
        p = _new_par(indent_chars=200, before=2)
        r = OxmlElement('w:r')
        _set_run(r, text, color=BLUE)
    else:
        p = _new_par(indent_chars=200, before=0)
        r = OxmlElement('w:r')
        _set_run(r, text)
    p.append(r)
    cell._tc.append(p)


def row_index(tbl, key, required=True):
    for i, tr in enumerate(tbl.findall(qn('w:tr'))):
        txt = ''.join(t.text or '' for t in tr.iter(qn('w:t')))
        if key in txt:
            return i
    if required:
        raise KeyError('表格中找不到行：%s' % key)
    return -1


def set_cell_plain(cell, text, color=None, bold=None):
    """表头字段：沿用原段落格式，只替换文字。

    入参既可能是 python-docx 的 _Cell，也可能已经是 lxml 的 w:tc 元素。
    """
    tc = getattr(cell, '_tc', cell)
    ps = tc.findall(qn('w:p'))
    for extra in ps[1:]:
        tc.remove(extra)
    p = ps[0]
    for r in list(p.findall(qn('w:r'))):
        p.remove(r)
    r = OxmlElement('w:r')
    _set_run(r, text, bold=bool(bold), color=color, size=24)
    p.append(r)


# ==================== 正文内容（与开题报告草稿.md 一致） ====================
# ==================== 正文内容 ====================
# 每个小节的内容组织为「段落列表」，一段表达一个完整意思，避免大段堆叠。
SEL = [
    '城市居住小区类型多样，商品房小区、长租公寓、公共租赁住房、单位与园区宿舍等均由相应的'
    '物业企业、运营单位或后勤部门实施管理。不同类型小区的管理主体不同，但面临的问题高度相似：'
    '目前多数小区的管理方仍以纸质台账、电子表格、电话和微信群等分散方式记录房屋信息、居住'
    '人员、报修申请和公共资源使用情况，存在信息更新滞后、居住关系不清晰、服务处理过程不可'
    '追溯、资源预约冲突频发、服务质量缺乏数据支撑等问题。这些问题的共同根源是信息分散在'
    '多个渠道，缺少一套贯穿居住关系全过程的统一记录，导致同一事实在不同环节重复录入、'
    '口径不一，出现争议时也难以回溯。',
    '本课题旨在设计并实现一个面向小区管理者的小区居住服务管理系统，基于 Spring Boot 与 Vue '
    '前后端分离架构，将小区房屋、居民居住关系、服务工单、公告反馈、公共资源预约及运营统计'
    '纳入统一管理，以规范小区日常服务流程、提高管理效率，并通过完整的设计与实现过程锻炼'
    '综合运用软件工程方法解决实际问题的能力。',
]

MEAN = [
    '本课题的实践意义集中在四个方面。一是将分散的小区管理信息集中化、结构化，通过入住审核、'
    '搬出登记和租期维护保证居住关系数据准确可查；二是将报修、隐患上报等事项统一为工单流程，'
    '处理全程留痕，居民可实时查看进度，管理方可以量化考核服务质量；三是通过公共资源预约的'
    '时间冲突校验避免资源重复占用，把冲突从事后协调提前到提交时拦截；四是通过统计报表为'
    '管理方提供运营决策依据，使服务质量与服务力量配置具备数据支撑。',
    '本课题的理论与训练意义在于：完整实践需求分析、系统设计、数据库设计、前后端编码、'
    '系统测试的软件工程全流程，掌握 Spring Boot、Vue 3、MySQL、Redis 等主流开发技术，'
    '提高独立分析问题与解决问题的能力。',
]

GOAL = [
    '面向具有明确服务管理方的居住小区的管理者，设计并实现支持单／多小区管理的 Web 系统，'
    '在空间维度上支持小区、楼栋、单元、房屋四级层级信息管理。居民与居住关系采用审核式管理，'
    '支持产权人、承租人与共同居住人等关系的登记与变更，注册方式（管理员代建或自助注册加'
    '入住审核）由后台配置，入住、搬出与租期变更均需全过程登记留痕。',
    '服务事项按性质分流：报修类事项走工单渠道全过程异步跟踪，意见建议与投诉类事项走反馈'
    '渠道并以会话形式实时沟通。系统同时提供公告定向广播（含已读统计）与消息通知中心，'
    '站内长连接实时推送，短信与微信服务号渠道可插拔；支持公共资源预约与时间冲突校验；'
    '面向潜在租户与购房者提供房源信息展示与看房预约；并提供服务评价与运营统计。',
    '系统设置居民、服务人员、小区管理员、超级管理员四类角色，权限分全局级（仅超级'
    '管理员）与小区级（小区管理员限所绑定小区）两层，多小区数据按小区维度隔离授权，'
    '关键操作留操作日志。',
]

ANA = [
    '通过调研住宅小区管理业务，抽象出四类角色与核心业务用例。未建立居住关系的居民仅可浏览公开公告'
    '与房源浏览、看房预约以及联系入口；居民侧覆盖查看居住信息与公告、提交服务申请与反馈、'
    '预约公共资源、参与反馈会话、确认与评价服务；服务人员侧聚焦工单的接收、处理与进度反馈，'
    '并可上传图片与文字处理结果。',
    '管理侧分两个层级。小区管理员面向所绑定的小区，负责基础信息维护、居住关系审核、工单'
    '派发、公告广播与反馈会话、资源配置、房源维护、通知渠道管理、本小区统计与居民账号管理；'
    '超级管理员面向全局，负责小区的新建与停用、全局配置、小区管理员账号创建与小区绑定、'
    '全局统计与操作日志。两级权限的分界是本系统权限设计的核心约束之一。',
    '业务渠道划分为双渠道：报修类事项走工单渠道，以异步流转的方式全过程跟踪；答复类事项'
    '走反馈渠道，以会话形式进行实时沟通。两条渠道共用同一套通知与留痕机制，但在状态模型'
    '与处理角色上彼此独立。对租住、工单、公共资源预约、看房预约四类业务分别建立状态模型，'
    '明确状态全集与合法流转。系统边界明确排除在线支付、房产交易环节与智能硬件接入等'
    '非管理核心功能。',
]

DES = [
    '系统采用前后端分离的单体架构，按「端—服务—数据」三层切分。前端以 Vue 3.5 与 TypeScript '
    '构建居民端、服务人员工作台与管理员后台三个入口，基于 Vite 构建、Pinia 管理状态、'
    'Axios 调用接口、ECharts 绘制统计图表。各端共用同一套请求封装与类型定义，仅在路由、'
    '布局与页面组织上按角色区分，既保证交互风格一致，又避免同一业务逻辑在前端重复实现。',
    '后端以 Spring Boot 4.0.6 提供 RESTful API，集成 Spring Security 与 JWT（jjwt）实现认证'
    '与基于角色的授权；多小区数据由 MyBatis-Plus 数据权限拦截器按小区维度统一过滤，业务'
    '代码中不重复编写可见范围判断。实时通信采用 WebSocket（SockJS + STOMP）承载反馈会话'
    '与消息通知推送，推送失败时由 HTTP 轮询增量补拉兜底，保证离线期间产生的消息不丢失。',
    '数据持久层采用 MySQL 9.x 存储业务数据，Redis 7 用于缓存与分布式锁（Redisson），ORM '
    '采用 MyBatis-Plus 3.5.16，表结构以迁移脚本为唯一事实来源，避免设计与实现出现两套口径；'
    '接口契约由 Springdoc 从代码注解生成，作为前后端联调的唯一依据。数据库设计覆盖小区、'
    '楼栋、单元、房屋、居民、居住关系、租住记录、服务类别、工单及处理记录、公告及查看记录、'
    '反馈会话、公共资源、预约、房源、看房预约、评价、通知与操作日志等核心表，绘制 E-R 图'
    '并遵循第三范式；对租住、工单、公共资源预约与看房预约四类业务分别设计状态机流转约束，'
    '明确状态全集、合法流转路径与越权流转的拒绝方式。',
]

IMP = [
    '系统实现遵循「分层落地、按域推进」的总体思路。后端按架构设计的分层契约组织为控制器层、'
    '服务层与数据访问层：控制器层只做参数校验与响应封装，业务规则集中在服务层，数据访问层'
    '由 ORM 框架承担且不承载业务判断。跨模块共用的实体、状态流转定义与通用存储集中在共享'
    '域内唯一维护，各端入口只负责本端的协调与查询，端与端之间不直接调用，公共层也不反向'
    '依赖端层，以此保证同一业务规则在多个入口下只有一份实现。',
    '实现顺序按业务依赖关系确定：先落地身份认证与权限，再依次实现小区基础信息、居住关系与'
    '租住、工单、公告与反馈会话、公共资源预约、房源与看房预约、消息通知中心，最后实现评价'
    '与统计。每个业务域按「实体—状态机—接口—页面」的顺序推进，前一域的状态与权限约束'
    '稳定后再进入下一域，避免因上游规则变动导致返工。',
    '实现过程中形成四条贯穿全系统的技术主线。其一是认证与授权主线：请求经统一的过滤器链'
    '完成身份解析，功能级权限按接口表达式判定，数据级权限由持久层拦截器按小区维度统一注入'
    '过滤条件，业务代码中不出现分散的可见范围判断。其二是状态流转主线：租住、工单、公共'
    '资源预约与看房预约四类业务各自维护状态全集与合法流转表，状态变更统一走状态机校验并'
    '同步写入处理记录，非法流转在提交前即被拒绝。其三是并发与一致性主线：预约类业务以时段'
    '记录加唯一约束表达占用，并发提交由数据库约束兜底、分布式锁降低无效重试；定时任务以'
    '防重入锁保证单实例执行，并以执行日志留痕。其四是通知与留痕主线：站内通知以长连接推送'
    '为主、增量补拉兜底，保证离线期间产生的消息不丢失；关键数据变更与协议确认均写入留痕'
    '记录，与业务数据在同一事务内提交。',
    '配套的定时任务承担租期状态自动判定、租期到期提醒、公告到期自动下线与浏览统计回写四类'
    '周期工作，任务之间互不阻塞，单次执行结果记入执行日志以便追溯。接口契约由接口文档工具'
    '从代码注解生成并作为前后端联调的唯一依据，前端各端按契约封装请求层与类型定义，减少'
    '联调阶段的语义偏差。',
]

TEST = [
    '测试以需求条目为唯一来源逐条推导：每一条功能需求与非功能需求都对应可执行的测试项，'
    '测试项再细化为具体用例，形成「需求—测试项—用例—执行结果—缺陷」的完整追溯链，'
    '任何一条需求都能回答「由哪些用例验证、结果如何」。测试在进入全面执行前先设冒烟环节，'
    '冒烟不通过则打回修改，不进入后续轮次，避免在基础功能不可用的状态下消耗测试资源。',
    '测试层次分为四类。功能测试按业务域组织，覆盖正常流程、边界条件与异常输入；权限测试'
    '采用否定用例，针对四类角色的越权访问与跨小区数据访问设计用例，要求越权请求全部被'
    '拒绝并留下记录；异常流程测试针对状态机的非法流转、重复提交与并发冲突构造场景；'
    '非功能测试覆盖接口响应时间、并发处理能力与长连接在弱网和断线重连下的表现。',
    '专项测试针对本系统中风险最高的三处设计单独建集。一是公共资源与看房预约的时间冲突'
    '校验，以重叠、相邻、跨天等时段组合验证唯一约束与合并提交的正确性，要求重叠预约全部'
    '被拦截；二是多小区数据隔离，以同一角色跨小区访问为场景验证数据级权限的过滤是否在'
    '所有查询路径上生效；三是通知的离线补偿，以「离线期间产生通知—重新上线—补拉可见」'
    '为主线验证消息不丢失且顺序正确。',
    '缺陷管理采用记录—修复—回归的闭环：每个缺陷记录复现条件、影响范围与等级，修复后必须'
    '在同一轮次内完成针对性回归，回归不通过不计入修复完成。测试出口以用例执行率、缺陷'
    '收敛情况与专项测试结果共同判定，测试结论与遗留问题一并写入测试报告，作为验收依据。',
]

KEYINTRO = '本课题的关键技术集中在权限、并发、实时通信与可视化四个方面，具体如下：'

KEYS = [
    '（1）基于 JWT 的无状态登录与接口级权限校验，配合功能权限与多小区数据权限两层控制；',
    '（2）基于 Redis 与 Redisson 的分布式锁，用于预约并发控制与定时任务防重入；',
    '（3）公共资源与看房预约的时间冲突检测，复用同一校验逻辑，重叠预约全部拦截；',
    '（4）租期到期状态的日期自动判定与到期提醒，以去重表保证同一租约不重复提醒；',
    '（5）WebSocket（SockJS + STOMP）实时会话与通知推送，结合增量游标补拉保证离线消息不丢失；',
    '（6）基于 MyBatis-Plus 数据权限拦截器实现小区级数据隔离；',
    '（7）操作日志切面、附件上传与 ECharts 统计可视化。',
]

METHOD = [
    '采用文献调查法、原型法与软件工程规范化方法相结合。通过文献调查明确小区居住服务管理的'
    '业务构成与前后端分离架构下的通用设计模式；通过原型法验证预约冲突、工单流转等关键交互。'
    '按「需求分析—系统设计—编码实现—测试验证—部署运行」的阶段主线推进，每一阶段产出'
    '对应文档，包括需求规格、设计文档与测试报告，保证开发过程可追溯、结论可复核。文档同时'
    '作为阶段之间的交付契约，下游阶段只依据上游文档推导任务，避免需求在设计或实现过程中'
    '被隐式改写。',
    '技术路线为：Java 17 与 Spring Boot 4.0.6 构成后端主体（Spring MVC、Spring Security、'
    'MyBatis-Plus、JWT、Redis、Redisson），Vue 3.5 与 TypeScript 构成前端主体（Vite、Pinia、'
    'Axios、ECharts），MySQL 9.x 作为数据库；前后端通过 RESTful API 交互，Maven 管理依赖，'
    'Flyway 迁移脚本作为表结构唯一事实来源。',
]

OUTLINE = [
    '1 绪论',
    '　1.1 研究背景及意义',
    '　1.2 国内外研究现状',
    '　1.3 研究内容与论文结构',
    '2 相关技术介绍',
    '　2.1 Spring Boot 框架',
    '　2.2 Vue 3 框架',
    '　2.3 MyBatis-Plus 与 Redis',
    '　2.4 系统开发运行环境',
    '3 系统分析',
    '　3.1 任务概述',
    '　3.2 可行性分析',
    '　3.3 功能需求分析（四类角色）',
    '　3.4 非功能需求分析',
    '4 系统设计',
    '　4.1 系统架构与功能模块设计',
    '　4.2 核心业务状态机设计（工单／租住／预约／看房预约）',
    '　4.3 数据库设计（设计原则／E-R 图／数据表设计）',
    '5 系统实现',
    '　5.1 用户认证与权限模块实现',
    '　5.2 小区基础信息与居住关系模块实现',
    '　5.3 工单与公告反馈会话模块实现',
    '　5.4 公共资源预约与房源看房模块实现',
    '　5.5 消息通知与统计模块实现',
    '6 系统测试',
    '　6.1 测试目的与方法',
    '　6.2 功能测试',
    '　6.3 权限与冲突校验专项测试',
    '　6.4 测试结论',
    '7 总结与展望',
]

HAVE = [
    '课程基础：已完成 Java 程序设计、数据库原理、软件工程、Web 开发等课程的学习，掌握 '
    'Spring Boot 与 Vue 的基础开发能力，并独立完成过课程级的前后端分离项目，对分层开发、'
    '接口联调与数据库设计有完整实践。',
    '技术栈验证：已完成 Java 17、Spring Boot 4.0.6、MyBatis-Plus 3.5.16、MySQL 9.x、'
    'Redis 7、Redisson、JWT、Vue 3.5、TypeScript、Vite 与 ECharts 的版本调研与验证；'
    '开发环境具备 JDK 17、Maven、Node.js、MySQL、Redis 等完整运行条件，可支撑开发、测试'
    '与部署的全过程。',
    '文档与定义：已建立项目文档体系并完成系统定义，明确目标用户、核心业务功能与验收判据，'
    '以及非目标边界；需求分析与系统设计的文档整理亦已完成，具备开展后续开发与论文撰写的'
    '条件。',
]

LACK = [
    '业务经验方面，对小区管理业务的实际流程了解主要来自调研资料与文献，尚缺少一线管理方的'
    '直接经验，对部分线下环节的处理惯例把握不够准确。',
    '工程实践方面，对 Spring Security 细粒度权限控制与预约冲突并发校验的实践尚不深入，'
    '在权限模型落地与高并发下的正确性验证上仍需积累经验。',
]

SOLVE = [
    '针对业务经验不足，向指导教师请教，并参考开源物业管理系统与相关文献，进一步校准业务模型。'
    '外文文献方面，已通过 OpenAlex 定向检索取得 9 篇全文（IEEE、MDPI、Springer、Elsevier '
    '期刊），后续如需补充可继续使用校园网数据库。',
    '针对权限与并发方面的不足，通过官方文档、开源案例与在线课程学习，并在系统测试阶段以'
    '专项用例集验证其正确性——权限方面以否定用例覆盖四类角色越权与跨小区访问，并发方面以'
    '重叠时段、重复提交等场景验证唯一约束与分布式锁的实际效果。',
]

# ==================== 生成 ====================
refs = extract_refs()
doc = Document(SRC)
t = doc.tables[0]
tbl = t._tbl


def set_header_paragraph(p, label, value, color=None):
    """表头段落（学院：… 专业班级：…）保留原有标签，只换值。"""
    for r in list(p._p.findall(qn('w:r'))):
        p._p.remove(r)
    r1 = OxmlElement('w:r')
    _set_run(r1, label, bold=False)
    r2 = OxmlElement('w:r')
    _set_run(r2, value, bold=True, color=color)
    p._p.append(r1)
    p._p.append(r2)


# ---------- 文档抬头：学院 / 专业班级 ----------
for p in doc.paragraphs:
    txt = p.text
    if '专业班级' in txt and '学院' in txt:
        set_header_paragraph(p, '学院：' + COLLEGE + '　　专业班级：', CLASS_NAME)
        break

# ---------- 表头字段（第 0 行：学生姓名 / 学号） ----------
set_cell_plain(t.rows[0].cells[1], STUDENT_NAME)
set_cell_plain(t.rows[0].cells[3], STUDENT_NO)

# ---------- 第 1 行：指导教师 ----------
set_cell_plain(t.rows[1].cells[1], ADVISOR)
# 该行应为 2 列：标签列 + 跨 3 列的值列
tr1 = tbl.findall(qn('w:tr'))[1]
tcs = tr1.findall(qn('w:tc'))
if len(tcs) == 4:
    for extra in tcs[2:]:
        tr1.remove(extra)
    tcPr = tcs[1].find(qn('w:tcPr'))
    if tcPr is not None:
        for gs in tcPr.findall(qn('w:gridSpan')):
            tcPr.remove(gs)
        gs = OxmlElement('w:gridSpan')
        gs.set(qn('w:val'), '3')
        tcPr.append(gs)

# ---------- 企业教师行（独立一行，置于指导教师之后） ----------
i_teacher = row_index(tbl, '指导教师')
tr_teacher = tbl.findall(qn('w:tr'))[i_teacher]

if row_index(tbl, '企业教师', required=False) < 0:
    tr_ent = copy.deepcopy(tbl.findall(qn('w:tr'))[row_index(tbl, '题目')])
    etcs = tr_ent.findall(qn('w:tc'))
    for extra in etcs[2:]:
        tr_ent.remove(extra)
    tc0 = etcs[0]
    for p0 in list(tc0.findall(qn('w:p'))):
        tc0.remove(p0)
    np0 = OxmlElement('w:p')
    ppr0 = OxmlElement('w:pPr')
    ps0 = OxmlElement('w:pStyle')
    ps0.set(qn('w:val'), '8')
    ppr0.append(ps0)
    jc0 = OxmlElement('w:jc')
    jc0.set(qn('w:val'), 'center')
    ppr0.append(jc0)
    np0.append(ppr0)
    r0 = OxmlElement('w:r')
    _set_run(r0, '企业教师')
    np0.append(r0)
    tc0.append(np0)
    set_cell_plain(etcs[1], ENTERPRISE_TEACHER)
    tr_teacher.addnext(tr_ent)

# 题目：定位原始「题目」行（位置可能已被移动，重新按文本查）
i_title = row_index(tbl, '题目')
set_cell_plain(tbl.findall(qn('w:tr'))[i_title].findall(qn('w:tc'))[1], TITLE)

# ---------- 勾选项：论文性质 / 选题类型 / 课题完成过程 ----------
def rebuild_checks(cell, items, mark_index, columns=(42, 42, 30)):
    """重建勾选行。

    items     : [文字]；mark_index 指定勾选第几项（0 基）
    columns   : 每一列的宽度百分比，长度即每行的项数

    复选框字符必须用 Wingdings 2 字体，否则 Word 中显示为缺字方框。
    每项之间用制表位分栏（而不是堆空格），满一行插入 w:br 换行，
    这样在 WPS / Word 中都不会出现「□理论研究型□应用研究型」粘连。
    """
    ps = cell._tc.findall(qn('w:p'))
    for extra in ps[1:]:
        cell._tc.remove(extra)
    p = ps[0]
    for r in list(p.findall(qn('w:r'))):
        p.remove(r)

    ppr = p.find(qn('w:pPr'))
    if ppr is None:
        ppr = OxmlElement('w:pPr')
        p.insert(0, ppr)
    for tag in ('w:ind', 'w:tabs'):
        old = ppr.find(qn(tag))
        if old is not None:
            ppr.remove(old)

    # 按列宽设置制表位（以缇为单位，A4 表格正文宽度约 8100 缇）
    tabs = OxmlElement('w:tabs')
    pos = 0
    for w in columns[:-1]:
        pos += int(8100 * w / 100.0)
        tab = OxmlElement('w:tab')
        tab.set(qn('w:val'), 'left')
        tab.set(qn('w:pos'), str(pos))
        tabs.append(tab)
    ppr.append(tabs)

    per_line = len(columns)
    for idx, label in enumerate(items):
        checked = (idx == mark_index)
        r = OxmlElement('w:r')
        rpr = OxmlElement('w:rPr')
        rf = OxmlElement('w:rFonts')
        for attr in ('w:ascii', 'w:hAnsi', 'w:eastAsia'):
            rf.set(qn(attr), 'Wingdings 2')
        rpr.append(rf)
        sz = OxmlElement('w:sz')
        sz.set(qn('w:val'), '24')
        rpr.append(sz)
        r.append(rpr)
        tt = OxmlElement('w:t')
        tt.set(qn('xml:space'), 'preserve')
        tt.text = '\uf052' if checked else '\uf0a3'
        r.append(tt)
        p.append(r)

        r2 = OxmlElement('w:r')
        _set_run(r2, label, bold=checked)
        p.append(r2)

        col = idx % per_line
        if idx == len(items) - 1:
            break
        if col == per_line - 1:
            # 换行符必须放在 run 内部
            br_run = OxmlElement('w:r')
            br = OxmlElement('w:br')
            br_run.append(br)
            p.append(br_run)
        else:
            tab_run = OxmlElement('w:r')
            tab_el = OxmlElement('w:tab')
            tab_run.append(tab_el)
            p.append(tab_run)


rebuild_checks(t.rows[row_index(tbl, '论文性质')].cells[1],
               ['毕业设计', '毕业论文'], 0, columns=(50, 50))
rebuild_checks(t.rows[row_index(tbl, '选题类型')].cells[1],
               ['软件设计型', '工程设计型', '艺术设计型',
                '理论研究型', '应用研究型', '其他'], 0, columns=(34, 34, 32))
# 模板没有「论文课题是否在以下过程中完成」行，按官方附件 1 补上
if row_index(tbl, '论文课题是否在以下', required=False) < 0:
    i_type = row_index(tbl, '选题类型')
    tr_type = tbl.findall(qn('w:tr'))[i_type]
    # 以「题目」行为模板（标签列 + gridSpan=3 的值列），不再依赖已被重建的「指导教师」行
    tr_done = copy.deepcopy(tbl.findall(qn('w:tr'))[row_index(tbl, '题目')])
    dcs = tr_done.findall(qn('w:tc'))
    for extra in dcs[2:]:
        tr_done.remove(extra)
    tc0 = dcs[0]
    for p in list(tc0.findall(qn('w:p'))):
        tc0.remove(p)
    p0 = OxmlElement('w:p')
    ppr0 = OxmlElement('w:pPr')
    ps0 = OxmlElement('w:pStyle')
    ps0.set(qn('w:val'), '8')
    ppr0.append(ps0)
    p0.append(ppr0)
    for line in ('论文课题是否在', '以下过程中完成'):
        r0 = OxmlElement('w:r')
        _set_run(r0, line)
        p0.append(r0)
        p0.append(OxmlElement('w:br'))
    tc0.append(p0)
    tr_type.addnext(tr_done)

rebuild_checks(t.rows[row_index(tbl, '论文课题是否在以下')].cells[1],
               ['工程实践', '实验、实习', '社会调查', '理论研究', '否'], 0,
               columns=(21, 21, 21, 21, 16))

# ---------- 一、选题的目的和意义 ----------
fill_cell(t.rows[row_index(tbl, '选题的目的')].cells[0], [
    ('一、选题的目的和意义', 'title'),
    ('1.1 选题目的', 'h2'), (SEL, 'body'),
    ('1.2 选题意义', 'h2'), (MEAN, 'body'),
])

# ---------- 二、主要研究内容 ----------
# 编号体系：一、二、三（章）→ 1.1 / 2.1（节）→ 2.1.1（条）
lines2 = [('二、主要研究内容（设计方案或论文撰写提纲）', 'title'),
          ('2.1 设计目标', 'h2'), (GOAL, 'body'),
          ('2.2 设计方案', 'h2'),
          ('2.2.1 系统分析', 'h3'), (ANA, 'body'),
          ('2.2.2 系统设计', 'h3'), (DES, 'body'),
          ('2.2.3 系统实现', 'h3'), (IMP, 'body'),
          ('2.2.4 系统测试', 'h3'), (TEST, 'body'),
          ('2.3 关键技术', 'h2'), (KEYINTRO, 'body')]
lines2 += [(k, 'body') for k in KEYS]
lines2 += [('2.4 拟采取的研究方法、技术路线', 'h2'), (METHOD, 'body'),
           ('2.5 论文撰写提纲', 'h2')]
lines2 += [(o, 'outline') for o in OUTLINE]
fill_cell(t.rows[row_index(tbl, '主要研究内容')].cells[0], lines2)

# ---------- 三、总体进度安排 ----------
fill_cell(t.rows[row_index(tbl, '总体进度安排')].cells[0], [
    ('三、总体进度安排（选题，开题、一稿、二稿、三稿、定稿等）', 'title'),
    ('2026年07月01日 至 2026年08月31日　　选题', 'outline'),
    ('2026年09月01日 至 2026年09月15日　　开题', 'outline'),
    ('2026年09月16日 至 2026年10月15日　　一稿', 'outline'),
    ('2026年10月16日 至 2026年10月31日　　二稿', 'outline'),
    ('2026年11月01日 至 2026年11月30日　　三稿', 'outline'),
    ('2026年12月01日 至 2026年12月30日　　定稿', 'outline'),
])

# ---------- 四、选题研究准备情况 ----------
n_cn = sum(1 for r in refs if not re.match(r'^[A-Z][A-Z\s,.]{3,}', r))
n_fo = len(refs) - n_cn
n_cn = sum(1 for r in refs if not re.match(r'^[A-Z][A-Z\s,.]{3,}', r))
n_fo = len(refs) - n_cn
fill_cell(t.rows[row_index(tbl, '选题研究准备情况')].cells[0],
          [('四、选题研究准备情况（已查阅主要参考文献，完成任务所具备的条件等）', 'title'),
           ('4.1 主要参考文献', 'h2'),
           ('本课题共参考 %d 篇文献，其中中文 %d 篇、外文 %d 篇，全部为 2021 年及以后'
            '发表，格式遵循 GB/T 7714-2015。' % (len(refs), n_cn, n_fo), 'body')]
          + [('[%d] %s' % (i + 1, r), 'ref') for i, r in enumerate(refs)]
          + [('4.2 已具备的研究条件', 'h2'), (HAVE, 'body'),
             ('4.3 尚缺少的研究条件', 'h2'), (LACK, 'body'),
             ('4.4 拟解决的途径', 'h2'), (SOLVE, 'body')])

# ---------- 五、指导教师意见 / 六、学院负责人审核意见 ----------
fill_cell(t.rows[row_index(tbl, '指导教师意见')].cells[0], [
    ('五、指导教师意见', 'title'),
    ('　　　　　　　　　　　　　　　　　　　　　　　　　　　　　　　　', 'body'),
    ('　　　　　　　　　　　　　　　　　　　　　　　　　　　　　　　　', 'body'),
    ('　　　　　　　　　　　　　　　　　　　　　　　　　　　　　　　　', 'body'),
    ('指导教师（签名）：　　　　　　　　　　　　年　　月　　日', 'body'),
])
fill_cell(t.rows[row_index(tbl, '学院负责人审核意见')].cells[0], [
    ('六、学院负责人审核意见（是否同意开题）', 'title'),
    ('　　　　　　　　　　　　　　　　　　　　　　　　　　　　　　　　', 'body'),
    ('　　　　　　　　　　　　　　　　　　　　　　　　　　　　　　　　', 'body'),
    ('学院负责人（签名）：　　　　　　　　　　　　年　　月　　日', 'body'),
])

# ---------- 页边距：收紧上下边距，避免表格跨页只漏一行 ----------
for _sec in doc.sections:
    _sec.top_margin = Cm(2.0)
    _sec.bottom_margin = Cm(2.0)

# ---------- 按大章节拆分表格并分页 ----------
# 整表跨页时 WPS 会把行内容裁掉，且单元格内的分页符不生效；
# 因此把每个大章节拆成独立表格，表格之间插入正文级分页符。
def split_row_before_text(tbl, row_idx, prefix):
    """把第 row_idx 行按段落拆成两行：从以 prefix 开头的段落起，移到新行。"""
    import copy
    trs = tbl.findall(qn('w:tr'))
    tr = trs[row_idx]
    tcs = tr.findall(qn('w:tc'))
    tc = tcs[0]
    ps = tc.findall(qn('w:p'))

    cut = None
    for i, para in enumerate(ps):
        txt = ''.join(x.text or '' for x in para.iter(qn('w:t'))).strip()
        if txt.startswith(prefix):
            cut = i
            break
    if cut is None or cut == 0:
        return False

    new_tr = copy.deepcopy(tr)
    new_tc = new_tr.findall(qn('w:tc'))[0]
    for p0 in new_tc.findall(qn('w:p')):
        new_tc.remove(p0)
    for para in ps[cut:]:
        tc.remove(para)
        new_tc.append(para)
    tr.addnext(new_tr)
    return True


def split_table_by_sections(tbl, groups):
    import copy
    all_rows = tbl.findall(qn('w:tr'))
    parent = tbl.getparent()
    pos = list(parent).index(tbl)
    parent.remove(tbl)
    for gi, (a, b) in enumerate(groups):
        new_tbl = copy.deepcopy(tbl)
        for tr in new_tbl.findall(qn('w:tr')):
            new_tbl.remove(tr)
        for tr in all_rows[a:b + 1]:
            new_tbl.append(copy.deepcopy(tr))
        parent.insert(pos, new_tbl)
        pos += 1
        if gi < len(groups) - 1:
            pbreak = OxmlElement('w:p')
            run = OxmlElement('w:r')
            br = OxmlElement('w:br')
            br.set(qn('w:type'), 'page')
            run.append(br)
            pbreak.append(run)
            parent.insert(pos, pbreak)
            pos += 1


# 先把「2.5 论文撰写提纲」从「二、」那一行拆出，避免提纲被页边界切断
_i2 = row_index(tbl, '二、主要研究内容')
split_row_before_text(tbl, _i2, '2.5 论文撰写提纲')

# 按标题文本重新定位（前面增删过表头行，不能用固定行号）
_i1 = row_index(tbl, '一、选题的目的')
_i2 = row_index(tbl, '二、主要研究内容')
_i25 = row_index(tbl, '2.5 论文撰写提纲')
_i3 = row_index(tbl, '三、总体进度安排')
_i4 = row_index(tbl, '四、选题研究准备情况')
_i5 = row_index(tbl, '五、指导教师意见')
_i6 = row_index(tbl, '六、学院负责人审核意见')

split_table_by_sections(tbl, [
    (0, _i1),          # 表头 + 一、
    (_i2, _i2),        # 二、（2.1~2.4）
    (_i25, _i25),      # 2.5 论文撰写提纲
    (_i3, _i3),        # 三、
    (_i4, _i4),        # 四、
    (_i5, _i6),        # 五、六、
])

if not _writable(OUT):
    OUT = OUT.replace('.docx', '_new.docx')
    print('目标文件被占用，改写为：', OUT)
doc.save(OUT)

# ---------- 清理模板携带的文档属性 ----------
doc2 = Document(OUT)
cp = doc2.core_properties
cp.author = STUDENT_NAME
cp.last_modified_by = STUDENT_NAME
cp.comments = ''
cp.category = ''
cp.subject = ''
cp.title = '本科毕业设计（论文）开题报告及任务书'
doc2.save(OUT)

print('生成完成：', OUT)
print('参考文献：%d 篇（中文 %d，外文 %d）' % (len(refs), n_cn, n_fo))
